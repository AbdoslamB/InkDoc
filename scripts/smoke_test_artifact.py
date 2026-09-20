#!/usr/bin/env python3
"""End-to-end smoke test runner for shipped InkDoc artifacts."""
from __future__ import annotations

import argparse
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path


def get_free_port() -> int:
    """Find a random ephemeral port that is free on 127.0.0.1."""
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def kill_process_tree(pid: int) -> None:
    """Forcefully terminate a process tree cross-platform."""
    if sys.platform == "win32":
        subprocess.run(
            ["taskkill", "/F", "/T", "/PID", str(pid)],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            check=False,
        )
    else:
        try:
            import signal

            pgid = os.getpgid(pid)
            os.killpg(pgid, signal.SIGTERM)
            time.sleep(0.5)
            os.killpg(pgid, signal.SIGKILL)
        except (ProcessLookupError, PermissionError):
            pass


def find_executable(base_dir: Path) -> Path:
    """Locate the InkDoc executable inside an extracted directory."""
    if sys.platform == "win32":
        candidates = list(base_dir.rglob("inkdoc.exe"))
        if not candidates:
            raise FileNotFoundError(f"Could not find inkdoc.exe in {base_dir}")
        return candidates[0]

    if sys.platform == "darwin":
        # Check for macOS .app bundle executable (Contents/MacOS/inkdoc)
        app_candidates = list(base_dir.rglob("Contents/MacOS/inkdoc"))
        if app_candidates:
            exe = app_candidates[0]
            exe.chmod(0o755)
            return exe
        raise FileNotFoundError(f"Could not find macOS .app bundle executable (Contents/MacOS/inkdoc) in {base_dir}")

    # Linux
    candidates = list(base_dir.rglob("inkdoc"))
    for c in candidates:
        if c.is_file() and not c.name.endswith(".zip"):
            c.chmod(0o755)
            return c
    raise FileNotFoundError(f"Could not find Linux inkdoc executable in {base_dir}")


def prepare_target_executable(artifact_path: Path, artifact_type: str, temp_dir: Path) -> Path:
    """Prepare executable path based on artifact type."""
    if artifact_type == "onefile":
        if not artifact_path.is_file():
            raise FileNotFoundError(f"Onefile artifact does not exist: {artifact_path}")
        return artifact_path

    if artifact_type == "zip":
        extract_dir = temp_dir / "unpacked"
        extract_dir.mkdir(parents=True, exist_ok=True)
        print(f"[*] Unpacking {artifact_path.name} to {extract_dir} ...")
        if sys.platform == "darwin":
            subprocess.run(["ditto", "-x", "-k", str(artifact_path), str(extract_dir)], check=True)
        elif shutil.which("unzip"):
            subprocess.run(["unzip", "-q", str(artifact_path), "-d", str(extract_dir)], check=True)
        else:
            with zipfile.ZipFile(artifact_path, "r") as z:
                z.extractall(extract_dir)
        return find_executable(extract_dir)

    if artifact_type == "installer":
        if sys.platform != "win32":
            raise RuntimeError("Installer smoke testing is only supported on Windows runners.")
        install_dir = temp_dir / "installed"
        install_dir.mkdir(parents=True, exist_ok=True)
        print(f"[*] Running silent Inno Setup installation to {install_dir} ...")
        cmd = [
            str(artifact_path),
            "/VERYSILENT",
            "/SUPPRESSMSGBOXES",
            "/NORESTART",
            f"/DIR={install_dir}",
        ]
        res = subprocess.run(cmd, capture_output=True, text=True)
        if res.returncode != 0:
            raise RuntimeError(
                f"Inno Setup installer failed with code {res.returncode}:\nSTDOUT: {res.stdout}\nSTDERR: {res.stderr}"
            )
        return find_executable(install_dir)

    raise ValueError(f"Unknown artifact type '{artifact_type}'. Expected: onefile, zip, or installer.")


def generate_sample_docx() -> bytes:
    """Generate minimal valid .docx file."""
    import io

    import docx

    doc = docx.Document()
    doc.add_heading("DOCX Smoke Test", level=1)
    doc.add_paragraph("Smoke test paragraph content for docx conversion.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_sample_xlsx() -> bytes:
    """Generate minimal valid .xlsx file."""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Metric", "Value", "Status"])
    ws.append(["TestRun", 42, "OK"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_sample_pdf() -> bytes:
    """Generate minimal valid PDF with extractable text using fpdf2."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(text="PDF Smoke Test Sample Document")
    return bytes(pdf.output())


def send_file_conversion(port: int, filename: str, content: bytes, content_type: str) -> str:
    """Test file conversion route via multipart upload and assert non-empty markdown."""
    boundary = "----SmokeTestBoundary123456789"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        f"Content-Type: {content_type}\r\n\r\n"
    ).encode() + content + f"\r\n--{boundary}--\r\n".encode()

    url = f"http://127.0.0.1:{port}/convert/file?save_to_downloads=false&response_format=json"
    req = urllib.request.Request(
        url,
        data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=20.0) as resp:
        if resp.getcode() != 200:
            raise RuntimeError(f"/convert/file for {filename} returned status {resp.getcode()}")
        data = json.loads(resp.read().decode("utf-8"))
        if not data.get("success"):
            raise RuntimeError(f"Conversion of {filename} failed: {data}")
        markdown = data.get("markdown", "")
        if not markdown.strip():
            raise RuntimeError(f"Conversion of {filename} returned empty markdown output!")
        return markdown


def run_smoke_test(
    artifact_path: Path,
    artifact_type: str,
    expected_version: str,
    timeout: float = 60.0,
    log_dir: Path | None = None,
) -> None:
    """Execute end-to-end smoke test on the specified artifact."""
    if not artifact_path.is_file():
        raise FileNotFoundError(f"Target artifact does not exist: {artifact_path}")

    log_dir = log_dir or (Path.cwd() / "smoke-logs")
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / f"{artifact_path.stem}.smoke.log"

    print(f"\n{'='*70}")
    print(f"SMOKE TEST: {artifact_path.name}")
    print(f"Type: {artifact_type} | Expected Version: {expected_version}")
    print(f"Log:  {log_file}")
    print(f"{'='*70}")

    with tempfile.TemporaryDirectory(prefix="inkdoc-smoke-", ignore_cleanup_errors=True) as tmp:
        temp_dir = Path(tmp)
        exe_path = prepare_target_executable(artifact_path, artifact_type, temp_dir)
        print(f"[*] Prepared executable: {exe_path}")

        # 1. Run GUI backend self-test in frozen build
        print(f"[*] Running GUI backend self-test: {exe_path} --selftest")
        selftest_proc = subprocess.Popen(
            [str(exe_path), "--selftest"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        try:
            out, err = selftest_proc.communicate(timeout=15.0)
            if selftest_proc.returncode != 0:
                raise RuntimeError(
                    f"--selftest failed with code {selftest_proc.returncode}:\n"
                    f"STDOUT:\n{out}\n"
                    f"STDERR:\n{err}"
                )
            print(f"[PASS] GUI backend self-test succeeded:\n{out.strip()}")
        except subprocess.TimeoutExpired as e:
            kill_process_tree(selftest_proc.pid)
            raise RuntimeError("--selftest timed out after 15s (process hung or showed modal crash dialog)") from e

        port = get_free_port()
        print(f"[*] Selected ephemeral port: {port}")

        with open(log_file, "w", encoding="utf-8", errors="replace") as log_f:
            kwargs: dict = {
                "stdout": log_f,
                "stderr": subprocess.STDOUT,
            }
            if sys.platform == "win32":
                kwargs["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
            else:
                kwargs["start_new_session"] = True

            cmd = [str(exe_path), "--headless", "--port", str(port)]
            print(f"[*] Launching: {' '.join(cmd)}")
            proc = subprocess.Popen(cmd, **kwargs)

            start_time = time.time()
            health_ok = False
            reported_version = None

            try:
                # 2. Poll /health with fast-fail
                print(f"[*] Polling http://127.0.0.1:{port}/health (timeout={timeout}s) ...")
                while (time.time() - start_time) < timeout:
                    # Detect process having exited (poll return code instead of matching log text)
                    ret = proc.poll()
                    if ret is not None:
                        log_f.flush()
                        log_content = log_file.read_text(encoding="utf-8", errors="replace") if log_file.is_file() else ""
                        print(f"\n--- Process Log ({log_file}) ---\n{log_content}\n--- End Process Log ---", file=sys.stderr)
                        raise RuntimeError(
                            f"Process exited prematurely with code {ret} before /health was ready.\n"
                            f"Process Log:\n{log_content}"
                        )

                    try:
                        req = urllib.request.Request(f"http://127.0.0.1:{port}/health")
                        with urllib.request.urlopen(req, timeout=2.0) as resp:
                            if resp.getcode() == 200:
                                data = json.loads(resp.read().decode("utf-8"))
                                reported_version = data.get("version")
                                if data.get("status") == "ok":
                                    health_ok = True
                                    print(f"[PASS] /health responded HTTP 200 OK (version='{reported_version}')")
                                    break
                    except (urllib.error.URLError, ConnectionResetError, OSError):
                        pass

                    time.sleep(0.5)

                if not health_ok:
                    log_f.flush()
                    log_content = log_file.read_text(encoding="utf-8", errors="replace") if log_file.is_file() else ""
                    print(f"\n--- Process Log ({log_file}) ---\n{log_content}\n--- End Process Log ---", file=sys.stderr)
                    raise TimeoutError(
                        f"Timed out after {timeout}s waiting for /health to become available.\n"
                        f"Process Log:\n{log_content}"
                    )

                # 3. Verify reported version matches expected tag/dev version
                if reported_version != expected_version:
                    raise AssertionError(
                        f"Version mismatch in /health: expected '{expected_version}', got '{reported_version}'"
                    )
                print(f"[PASS] Version verification passed: '{reported_version}' == '{expected_version}'")

                # 4. Multi-format conversion testing
                print("[*] Testing document conversions via /convert/file...")

                # 4a. Plain text (.txt)
                txt_data = b"# Sample Heading\nSmoke test body content for InkDoc verification.\r\n"
                txt_md = send_file_conversion(port, "sample.txt", txt_data, "text/plain")
                if "Sample Heading" not in txt_md:
                    raise AssertionError(f"Expected content in text markdown output, got: {txt_md[:200]}")
                print(f"[PASS] Plain text (.txt) converted successfully ({len(txt_md)} bytes markdown).")

                # 4b. Word Document (.docx)
                docx_data = generate_sample_docx()
                docx_md = send_file_conversion(
                    port,
                    "sample.docx",
                    docx_data,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
                if "DOCX Smoke Test" not in docx_md:
                    raise AssertionError(f"Expected content in docx markdown output, got: {docx_md[:200]}")
                print(f"[PASS] Word document (.docx) converted successfully ({len(docx_md)} bytes markdown).")

                # 4c. Excel Workbook (.xlsx)
                xlsx_data = generate_sample_xlsx()
                xlsx_md = send_file_conversion(
                    port,
                    "sample.xlsx",
                    xlsx_data,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                if "TestRun" not in xlsx_md and "Metric" not in xlsx_md:
                    raise AssertionError(f"Expected content in xlsx markdown output, got: {xlsx_md[:200]}")
                print(f"[PASS] Excel workbook (.xlsx) converted successfully ({len(xlsx_md)} bytes markdown).")

                # 4d. PDF Document (.pdf)
                pdf_data = generate_sample_pdf()
                pdf_md = send_file_conversion(port, "sample.pdf", pdf_data, "application/pdf")
                if "PDF Smoke Test" not in pdf_md:
                    raise AssertionError(f"Expected content in pdf markdown output, got: {pdf_md[:200]}")
                print(f"[PASS] PDF document (.pdf) converted successfully ({len(pdf_md)} bytes markdown).")

            finally:
                print(f"[*] Terminating process tree for PID {proc.pid} ...")
                kill_process_tree(proc.pid)
                try:
                    proc.wait(timeout=5)
                except Exception:
                    pass
                time.sleep(1.0)

    print(f"[ALL PASS] Smoke test succeeded for {artifact_path.name}\n")


def main() -> int:
    parser = argparse.ArgumentParser(description="Smoke-test shipped InkDoc artifacts.")
    parser.add_argument("--artifact", type=Path, required=True, help="Path to artifact to test.")
    parser.add_argument(
        "--type",
        choices=["onefile", "zip", "installer"],
        required=True,
        help="Artifact distribution packaging type.",
    )
    parser.add_argument("--expected-version", required=True, help="Expected version reported by /health.")
    parser.add_argument("--timeout", type=float, default=60.0, help="Timeout in seconds for /health polling.")
    parser.add_argument("--log-dir", type=Path, default=None, help="Directory to save process logs.")

    args = parser.parse_args()

    try:
        run_smoke_test(
            artifact_path=args.artifact,
            artifact_type=args.type,
            expected_version=args.expected_version,
            timeout=args.timeout,
            log_dir=args.log_dir,
        )
        return 0
    except Exception as e:
        print(f"\n[FAIL] Smoke test failed: {e}", file=sys.stderr)
        log_file = (args.log_dir or (Path.cwd() / "smoke-logs")) / f"{args.artifact.stem}.smoke.log"
        if log_file.is_file():
            print(f"\n--- Process Log ({log_file}) ---", file=sys.stderr)
            print(log_file.read_text(encoding="utf-8", errors="replace"), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
